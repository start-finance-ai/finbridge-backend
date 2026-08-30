from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Sequence
from xml.etree.ElementTree import ParseError as XMLParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 50_000
SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}

REQUIRED_COLUMNS = ("transaction_date", "sales_amount")
OPTIONAL_COLUMNS = ("transaction_id", "category")
COLUMN_ALIASES = {
    "transaction_date": "transaction_date",
    "거래일자": "transaction_date",
    "매출일자": "transaction_date",
    "sales_amount": "sales_amount",
    "매출액": "sales_amount",
    "판매금액": "sales_amount",
    "transaction_id": "transaction_id",
    "거래ID": "transaction_id",
    "거래번호": "transaction_id",
    "category": "category",
    "카테고리": "category",
    "분류": "category",
}

_AMOUNT_PATTERN = re.compile(
    r"^-?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?$"
)
_XLSX_PARSE_EXCEPTIONS = (
    BadZipFile,
    InvalidFileException,
    XMLParseError,
    EOFError,
    KeyError,
    OSError,
    TypeError,
    ValueError,
)


class SalesUploadError(ValueError):
    def __init__(
        self,
        error_code: str,
        message: str,
        *,
        status_code: int = 422,
        row_number: int | None = None,
        field: str | None = None,
        reason: str | None = None,
        context: dict[str, Any] | None = None,
    ) -> None:
        super().__init__(message)
        self.error_code = error_code
        self.message = message
        self.status_code = status_code
        self.row_number = row_number
        self.field = field
        self.reason = reason
        self.context = context or {}

    def to_detail(self) -> dict[str, Any]:
        detail: dict[str, Any] = {
            "error_code": self.error_code,
            "message": self.message,
        }
        if self.row_number is not None:
            detail["row_number"] = self.row_number
        if self.field is not None:
            detail["field"] = self.field
        if self.reason is not None:
            detail["reason"] = self.reason
        detail.update(self.context)
        return detail


@dataclass(frozen=True)
class SalesRow:
    transaction_date: date
    sales_amount: Decimal


@dataclass(frozen=True)
class ParsedSalesUpload:
    rows: list[SalesRow]
    source_format: str
    sheet_name: str | None
    rows_received: int
    ignored_columns: list[str]


@dataclass(frozen=True)
class HeaderAnalysis:
    indices: dict[str, int]
    ignored_columns: list[str]
    missing_columns: list[str]
    conflicts: dict[str, list[str]]


def parse_sales_upload(*, filename: str | None, content: bytes) -> ParsedSalesUpload:
    extension = Path(filename or "").suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise SalesUploadError(
            "UNSUPPORTED_FILE_TYPE",
            "Only .csv and .xlsx files are supported",
            status_code=415,
            context={"supported_extensions": sorted(SUPPORTED_EXTENSIONS)},
        )
    if not content:
        raise SalesUploadError(
            "EMPTY_FILE",
            "Uploaded file is empty",
            status_code=400,
        )
    if len(content) > MAX_FILE_BYTES:
        raise SalesUploadError(
            "FILE_TOO_LARGE",
            "Uploaded file exceeds the 5 MB limit",
            status_code=413,
            context={"max_file_bytes": MAX_FILE_BYTES},
        )

    if extension == ".csv":
        return _parse_csv(content)
    return _parse_xlsx(content)


def _parse_csv(content: bytes) -> ParsedSalesUpload:
    text = _decode_csv(content)
    try:
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        header_row_number, headers = _next_non_empty_row(reader)
        if headers is None:
            raise SalesUploadError(
                "EMPTY_FILE",
                "CSV does not contain a header row",
                status_code=400,
            )
        header = _validate_headers(headers)
        rows, rows_received = _parse_rows(
            reader,
            header=header,
            starting_row_number=header_row_number + 1,
        )
    except SalesUploadError:
        raise
    except csv.Error as exc:
        raise SalesUploadError(
            "PARSE_ERROR",
            "CSV parsing failed",
            reason=str(exc),
        ) from None

    return ParsedSalesUpload(
        rows=rows,
        source_format="CSV",
        sheet_name=None,
        rows_received=rows_received,
        ignored_columns=header.ignored_columns,
    )


def _parse_xlsx(content: bytes) -> ParsedSalesUpload:
    try:
        workbook = load_workbook(
            filename=io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except _XLSX_PARSE_EXCEPTIONS:
        raise SalesUploadError(
            "PARSE_ERROR",
            "XLSX workbook could not be parsed",
        ) from None

    try:
        candidates: list[tuple[Any, int, HeaderAnalysis]] = []
        conflicting_candidates: list[tuple[str, HeaderAnalysis]] = []
        for worksheet in workbook.worksheets:
            header_row_number, headers = _worksheet_header(worksheet)
            if headers is None:
                continue
            analysis = _analyze_headers(headers)
            if not analysis.missing_columns and analysis.conflicts:
                conflicting_candidates.append((worksheet.title, analysis))
            elif not analysis.missing_columns:
                candidates.append((worksheet, header_row_number, analysis))

        if conflicting_candidates:
            sheet_name, analysis = conflicting_candidates[0]
            _raise_column_conflict(analysis, sheet_name=sheet_name)
        if not candidates:
            raise SalesUploadError(
                "REQUIRED_COLUMN_MISSING",
                "No worksheet contains all required columns",
                field=",".join(REQUIRED_COLUMNS),
                context={"sheet_names": workbook.sheetnames},
            )
        if len(candidates) > 1:
            raise SalesUploadError(
                "MULTIPLE_VALID_SHEETS",
                "More than one worksheet matches the FinBridge Sample Schema",
                context={
                    "sheet_names": [worksheet.title for worksheet, _, _ in candidates]
                },
            )

        worksheet, header_row_number, header = candidates[0]
        rows, rows_received = _parse_rows(
            (
                list(row)
                for row_number, row in enumerate(
                    worksheet.iter_rows(values_only=True), start=1
                )
                if row_number > header_row_number
            ),
            header=header,
            starting_row_number=header_row_number + 1,
        )
        return ParsedSalesUpload(
            rows=rows,
            source_format="XLSX",
            sheet_name=worksheet.title,
            rows_received=rows_received,
            ignored_columns=header.ignored_columns,
        )
    except SalesUploadError:
        raise
    except _XLSX_PARSE_EXCEPTIONS:
        raise SalesUploadError(
            "PARSE_ERROR",
            "XLSX workbook could not be parsed",
        ) from None
    finally:
        workbook.close()


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SalesUploadError(
        "PARSE_ERROR",
        "CSV encoding must be UTF-8, UTF-8-SIG, or CP949",
        reason="UNSUPPORTED_ENCODING",
    )


def _next_non_empty_row(
    rows: Iterable[Sequence[Any]],
) -> tuple[int, list[Any] | None]:
    for row_number, row in enumerate(rows, start=1):
        values = list(row)
        if not _is_blank_row(values):
            return row_number, values
    return 0, None


def _worksheet_header(worksheet: Any) -> tuple[int, list[Any] | None]:
    return _next_non_empty_row(worksheet.iter_rows(values_only=True))


def _analyze_headers(headers: Sequence[Any]) -> HeaderAnalysis:
    indices: dict[str, int] = {}
    resolved_names: dict[str, list[str]] = {}
    ignored_columns: list[str] = []

    for index, raw_header in enumerate(headers):
        name = str(raw_header).strip() if raw_header is not None else ""
        canonical = COLUMN_ALIASES.get(name)
        if canonical is None:
            if name and name not in ignored_columns:
                ignored_columns.append(name)
            continue
        resolved_names.setdefault(canonical, []).append(name)
        indices.setdefault(canonical, index)

    conflicts = {
        canonical: names
        for canonical, names in resolved_names.items()
        if len(names) > 1
    }
    missing_columns = [
        column for column in REQUIRED_COLUMNS if column not in indices
    ]
    return HeaderAnalysis(
        indices=indices,
        ignored_columns=ignored_columns,
        missing_columns=missing_columns,
        conflicts=conflicts,
    )


def _validate_headers(headers: Sequence[Any]) -> HeaderAnalysis:
    analysis = _analyze_headers(headers)
    if analysis.conflicts:
        _raise_column_conflict(analysis)
    if analysis.missing_columns:
        raise SalesUploadError(
            "REQUIRED_COLUMN_MISSING",
            "Required columns are missing",
            field=",".join(analysis.missing_columns),
            context={"missing_columns": analysis.missing_columns},
        )
    return analysis


def _raise_column_conflict(
    analysis: HeaderAnalysis, *, sheet_name: str | None = None
) -> None:
    canonical, columns = next(iter(analysis.conflicts.items()))
    context: dict[str, Any] = {"conflicting_columns": columns}
    if sheet_name is not None:
        context["sheet_name"] = sheet_name
    raise SalesUploadError(
        "COLUMN_CONFLICT",
        "Multiple columns resolve to the same canonical column",
        field=canonical,
        context=context,
    )


def _parse_rows(
    rows: Iterable[Sequence[Any]],
    *,
    header: HeaderAnalysis,
    starting_row_number: int,
) -> tuple[list[SalesRow], int]:
    parsed_rows: list[SalesRow] = []
    rows_received = 0
    for row_number, row in enumerate(rows, start=starting_row_number):
        values = list(row)
        if _is_blank_row(values):
            continue
        rows_received += 1
        if rows_received > MAX_DATA_ROWS:
            raise SalesUploadError(
                "ROW_LIMIT_EXCEEDED",
                "Uploaded file exceeds the 50,000 data row limit",
                status_code=413,
                row_number=row_number,
                context={"max_data_rows": MAX_DATA_ROWS},
            )
        transaction_date = _parse_transaction_date(
            _value_at(values, header.indices["transaction_date"]),
            row_number=row_number,
        )
        sales_amount = _parse_sales_amount(
            _value_at(values, header.indices["sales_amount"]),
            row_number=row_number,
        )
        parsed_rows.append(
            SalesRow(
                transaction_date=transaction_date,
                sales_amount=sales_amount,
            )
        )

    if not parsed_rows:
        raise SalesUploadError(
            "EMPTY_DATA",
            "File does not contain any data rows",
        )
    return parsed_rows, rows_received


def _parse_transaction_date(value: Any, *, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        normalized = value.strip()
        for date_format in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(normalized, date_format).date()
            except ValueError:
                continue
    raise SalesUploadError(
        "INVALID_DATE",
        "transaction_date must use YYYY-MM-DD or YYYY/MM/DD",
        row_number=row_number,
        field="transaction_date",
        reason="UNSUPPORTED_OR_INVALID_DATE",
    )


def _parse_sales_amount(value: Any, *, row_number: int) -> Decimal:
    if isinstance(value, bool) or value is None:
        _raise_invalid_amount(row_number, "EMPTY_OR_NON_NUMERIC")

    try:
        if isinstance(value, Decimal):
            amount = value
        elif isinstance(value, int):
            amount = Decimal(value)
        elif isinstance(value, float):
            amount = Decimal(str(value))
        elif isinstance(value, str):
            normalized = value.strip()
            if not normalized or _AMOUNT_PATTERN.fullmatch(normalized) is None:
                _raise_invalid_amount(row_number, "UNSUPPORTED_AMOUNT_FORMAT")
            amount = Decimal(normalized.replace(",", ""))
        else:
            _raise_invalid_amount(row_number, "NON_NUMERIC")
    except InvalidOperation:
        _raise_invalid_amount(row_number, "INVALID_DECIMAL")

    if not amount.is_finite():
        _raise_invalid_amount(row_number, "NON_FINITE")
    if amount < 0:
        raise SalesUploadError(
            "NEGATIVE_SALES_AMOUNT",
            "Negative sales_amount is not supported",
            row_number=row_number,
            field="sales_amount",
            reason="REFUND_OR_CANCELLATION_NOT_SUPPORTED",
        )
    return amount


def _raise_invalid_amount(row_number: int, reason: str) -> None:
    raise SalesUploadError(
        "INVALID_SALES_AMOUNT",
        "sales_amount must be a finite non-negative number",
        row_number=row_number,
        field="sales_amount",
        reason=reason,
    )


def _is_blank_row(values: Sequence[Any]) -> bool:
    return all(
        value is None or (isinstance(value, str) and not value.strip())
        for value in values
    )


def _value_at(values: Sequence[Any], index: int) -> Any:
    return values[index] if index < len(values) else None
